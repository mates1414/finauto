"""Bi-directional Excel workflow (Phase 3, A4): read a user-corrected workbook
back into the Pydantic contract, diff it against the original, and surface the
computed outputs for the report stage.

The one hard gotcha (documented in plan.md / phase3.md): XlsxWriter ships
formulas **without cached values**, so the *computed* cells (target price, WACC,
multiples, signal) are blank until the workbook is recalculated. `recalc` does
that with headless LibreOffice (primary) or the `formulas` library (fallback);
`read_inputs` then reads with openpyxl `data_only=True`. The *input* (blue) cells
are stored literals and read back fine even without a recalc.

Everything here is pure except `recalc`, which shells out to LibreOffice or runs
the `formulas` calculator — the single network/binary edge in this module.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..schemas import Assumptions, CompanyFinancials, EditNote, FiscalYearData, Number
from .formulas import S01, S02, S06
from .sheets.s01_assumptions import NAME_CELLS
from .sheets.s02_historicals import ROWS as S02_ROWS


class RecalcError(Exception):
    pass


# s02 input rows -> (statement section, schema field). Mirrors the hardcoded
# `_VALUE_GETTERS` in s02_historicals: only reported (input) lines appear here;
# gross_profit / ebit / net_debt are formulas and are intentionally excluded.
_ROW_TO_FIELD: dict[str, tuple[str, str]] = {
    "revenue": ("income_statement", "revenue"),
    "cogs": ("income_statement", "cogs"),
    "sga": ("income_statement", "sga"),
    "ebitda": ("income_statement", "ebitda"),
    "da": ("income_statement", "depreciation_amortization"),
    "net_interest": ("income_statement", "net_interest_expense"),
    "net_income": ("income_statement", "net_income"),
    "cash": ("balance_sheet", "cash_and_equivalents"),
    "current_assets": ("balance_sheet", "total_current_assets"),
    "total_assets": ("balance_sheet", "total_assets"),
    "st_debt": ("balance_sheet", "short_term_debt"),
    "lt_debt": ("balance_sheet", "long_term_debt"),
    "leases": ("balance_sheet", "lease_liabilities"),
    "total_liabilities": ("balance_sheet", "total_liabilities"),
    "retained": ("balance_sheet", "retained_earnings"),
    "equity": ("balance_sheet", "total_equity"),
    "capex": ("cash_flow", "capex"),
}

# s01 defined-name cells -> Assumptions field. CurrentPrice/SharesOut are market
# inputs (not Assumptions) and are returned in the `computed` dict instead.
_NAME_TO_ASSUMPTION: dict[str, str] = {
    "RiskFree": "risk_free_rate",
    "ERP": "equity_risk_premium",
    "CRP": "country_risk_premium",
    "TaxRate": "tax_rate",
    "FXRate": "fx_rate",
    "Growth1": "growth_stage1",
    "Growth2": "growth_stage2",
    "TerminalGrowth": "terminal_growth",
    "EBITMargin": "ebit_margin",
    "CapExPct": "capex_pct_sales",
    "NWCPct": "nwc_pct_sales",
    "DAPct": "da_pct_sales",
    "WeightDCF": "weight_dcf",
    "WeightMultiples": "weight_multiples",
    "SignalThreshold": "signal_threshold",
}

# Computed outputs reachable via workbook-level defined names (set by the builder).
# A faithful recalc (Excel-on-open, headless LibreOffice) preserves these names.
_COMPUTED_NAMES: dict[str, str] = {
    "target_price": "TargetPrice",
    "dcf_price": "DCFImpliedPrice",
    "ev_ebitda_price": "PriceEVEBITDA",
    "ev_sales_price": "PriceEVSales",
    "pe_price": "PricePE",
    "wacc": "WACC",
    "cost_of_equity": "CostOfEquity",
    "cost_of_debt": "CostOfDebt",
    "current_price": "CurrentPrice",
    "shares_out": "SharesOut",
}

# Fixed-coordinate fallbacks for outputs whose cell is peer-count independent, used
# when defined names are absent (the `formulas` recalc fallback drops them). The
# 06_Valuation_Summary layout is static; WACC/Ke/Rd live on 03 at peer-dependent
# rows, so they resolve by defined name only and stay blank on the formulas path.
_COMPUTED_FALLBACK: dict[str, tuple[str, str]] = {
    "target_price": (S06, "B9"),
    "current_price": (S06, "B10"),
    "dcf_price": (S06, "B4"),
    "ev_ebitda_price": (S06, "B5"),
    "pe_price": (S06, "B6"),
    "ev_sales_price": (S06, "B7"),
    "shares_out": (S01, "B24"),
}


def recalc(path: Path | str, out_dir: Optional[Path] = None) -> Path:
    """Materialize formula values into a new workbook and return its path.

    Primary: headless LibreOffice (``soffice --convert-to xlsx``). Fallback: the
    ``formulas`` library. Raises if neither is available with a clear message.
    """
    src = Path(path)
    out_dir = (
        Path(out_dir) if out_dir else Path(tempfile.mkdtemp(prefix="finauto_recalc_"))
    )
    out_dir.mkdir(parents=True, exist_ok=True)

    lo = _recalc_with_libreoffice(src, out_dir)
    if lo is not None:
        return lo
    fm = _recalc_with_formulas(src, out_dir)
    if fm is not None:
        return fm
    raise RecalcError(
        "Cannot recalculate the workbook: neither LibreOffice (soffice) nor the "
        "`formulas` library is available. Install LibreOffice, or "
        "`pip install finauto[report]` for the `formulas` fallback."
    )


def _recalc_with_libreoffice(src: Path, out_dir: Path) -> Optional[Path]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    # Re-saving through LibreOffice recomputes formulas and stores cached values.
    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out_dir),
                str(src),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        raise RecalcError(f"LibreOffice recalc failed: {e}") from e
    out = out_dir / f"{src.stem}.xlsx"
    return out if out.exists() else None


def _recalc_with_formulas(src: Path, out_dir: Path) -> Optional[Path]:
    try:
        import formulas
    except ImportError:
        return None
    try:
        xl_model = formulas.ExcelModel().loads(str(src)).finish()
        xl_model.calculate()
        xl_model.write(dirpath=str(out_dir))
    except Exception as e:  # the lib raises various errors on unsupported funcs
        raise RecalcError(f"`formulas` recalc failed: {e}") from e
    # `formulas` writes <BOOKNAME>.xlsx (often upper-cased) into dirpath.
    candidates = list(out_dir.glob("*.xlsx"))
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def read_inputs(
    path: Path | str, ticker: str = "", name: Optional[str] = None
) -> tuple[CompanyFinancials, Assumptions, dict[str, object]]:
    """Read a (recalculated) workbook back into the Pydantic contract.

    Returns (financials in absolute units, assumptions, computed-outputs dict).
    The blue input cells are read regardless of recalc; computed outputs are
    blank until the workbook has been recalculated (see `recalc`).
    """
    wb = load_workbook(path, data_only=True)
    try:
        financials = _read_financials(wb, ticker=ticker, name=name)
        assumptions = _read_assumptions(wb)
        computed = _read_computed(wb)
    finally:
        wb.close()
    return financials, assumptions, computed


def _ws(wb, title: str):
    """Resolve a worksheet by title, case-insensitively (the `formulas` recalc
    fallback upper-cases sheet names; Excel/LibreOffice preserve them)."""
    if title in wb.sheetnames:
        return wb[title]
    low = title.lower()
    for name in wb.sheetnames:
        if name.lower() == low:
            return wb[name]
    raise KeyError(f"worksheet {title!r} not found (have {wb.sheetnames})")


def _num(value) -> Number:
    if isinstance(value, bool):
        return None
    return float(value) if isinstance(value, (int, float)) else None


def _read_financials(wb, *, ticker: str, name: Optional[str]) -> CompanyFinancials:
    ws = _ws(wb, S02)
    year_cols = _year_columns(ws)
    periods: list[FiscalYearData] = []
    for year, col0 in year_cols.items():
        column = get_column_letter(col0 + 1)
        period = FiscalYearData(year=year)
        for row_key, (section, field) in _ROW_TO_FIELD.items():
            value = ws[f"{column}{S02_ROWS[row_key] + 1}"].value
            if isinstance(value, (int, float)):
                setattr(getattr(period, section), field, float(value))
        periods.append(period)
    # Inputs are written to the sheet already normalized to absolute units.
    return CompanyFinancials(ticker=ticker, name=name, units="units", periods=periods)


def _year_columns(ws) -> dict[int, int]:
    """Map fiscal year -> 0-based column by scanning the s02 year header row."""
    header_row = 2  # 0-based; the year header row is fixed at row 3 (1-based) in s02
    cols: dict[int, int] = {}
    col0 = 1
    while True:
        value = ws.cell(row=header_row + 1, column=col0 + 1).value
        if not isinstance(value, int):
            break
        cols[value] = col0
        col0 += 1
    return cols


def _read_assumptions(wb) -> Assumptions:
    ws = _ws(wb, S01)
    overrides: dict[str, float] = {}
    for name_key, field in _NAME_TO_ASSUMPTION.items():
        coord = NAME_CELLS[name_key].replace("$", "")
        value = _num(ws[coord].value)
        if value is not None:
            overrides[field] = value
    return Assumptions(**overrides)


def _read_computed(wb) -> dict[str, object]:
    computed: dict[str, object] = {}
    for key, defined_name in _COMPUTED_NAMES.items():
        value = _resolve_name(wb, defined_name)
        if value is None and key in _COMPUTED_FALLBACK:
            sheet, coord = _COMPUTED_FALLBACK[key]
            value = _num(_ws(wb, sheet)[coord].value)
        computed[key] = value
    ws6 = _ws(wb, S06)
    computed["upside"] = _num(ws6["B11"].value)
    computed["fx_price"] = _num(ws6["B14"].value)
    # Signal is a text label (AL/TUT/SAT), not a number.
    signal = ws6["B12"].value
    computed["signal"] = signal if isinstance(signal, str) else None
    return computed


def read_signal(path: Path | str) -> Optional[str]:
    """Read the AL/TUT/SAT signal text from a recalculated workbook."""
    wb = load_workbook(path, data_only=True)
    try:
        value = _ws(wb, S06)["B12"].value
    finally:
        wb.close()
    return value if isinstance(value, str) else None


def _resolve_name(wb, defined_name: str) -> Number:
    dn = wb.defined_names.get(defined_name)
    if dn is None:
        return None
    for sheet_title, coord in dn.destinations:
        return _num(_ws(wb, sheet_title)[coord.replace("$", "")].value)
    return None


def _close(a: Number, b: Number, *, rel: float = 1e-6, abs_: float = 1e-6) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= max(abs_, rel * max(abs(a), abs(b)))


def diff_inputs(
    original: CompanyFinancials, edited: CompanyFinancials
) -> list[EditNote]:
    """Report every input cell the user changed (original vs edited financials).

    Both sides are normalized to absolute units and deduped by year before
    comparison, so the diff is unit- and ordering-independent.
    """
    orig = {p.year: p for p in original.normalized().deduped_periods()}
    edit = {p.year: p for p in edited.normalized().deduped_periods()}
    notes: list[EditNote] = []
    for year in sorted(set(orig) | set(edit)):
        o = orig.get(year)
        e = edit.get(year)
        for _row_key, (section, field) in _ROW_TO_FIELD.items():
            ov = getattr(getattr(o, section), field) if o else None
            ev = getattr(getattr(e, section), field) if e else None
            if not _close(ov, ev):
                notes.append(EditNote(path=f"{year}.{section}.{field}", old=ov, new=ev))
    return notes
