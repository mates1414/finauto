"""Read the generated workbook back with openpyxl and assert structure,
formula strings, defined names, and the no-unguarded-division invariant."""

from __future__ import annotations

import openpyxl
import pytest

SHEETS = [
    "01_Assumptions",
    "02_Historical_Financials",
    "03_WACC_Calculation",
    "04_DCF_Model",
    "05_Relative_Valuation",
    "06_Valuation_Summary",
]

DEFINED_NAMES = {
    "RiskFree", "ERP", "CRP", "TaxRate", "FXRate",
    "Growth1", "Growth2", "TerminalGrowth", "EBITMargin",
    "CapExPct", "NWCPct", "DAPct",
    "WeightDCF", "WeightMultiples", "SignalThreshold",
    "CurrentPrice", "SharesOut",
    "WACC", "CostOfEquity", "CostOfDebt",
    "DCFImpliedPrice", "PriceEVEBITDA", "PriceEVSales", "PricePE",
    "TargetPrice",
}


@pytest.fixture(scope="module")
def wb(workbook_path):
    return openpyxl.load_workbook(workbook_path)  # formulas kept as strings


def test_sheet_names(wb):
    assert wb.sheetnames == SHEETS


def test_defined_names_registered(wb):
    assert DEFINED_NAMES <= set(wb.defined_names)
    assert "03_WACC_Calculation" in wb.defined_names["WACC"].value
    assert wb.defined_names["RiskFree"].value == "'01_Assumptions'!$B$4"


def test_assumption_inputs_are_values(wb, inputs):
    ws = wb["01_Assumptions"]
    assert ws["B4"].value == pytest.approx(inputs.assumptions.risk_free_rate)
    assert ws["B7"].value == pytest.approx(inputs.assumptions.tax_rate)
    assert ws["B23"].value == pytest.approx(95.5)  # current price prefilled
    assert ws["B24"].value == pytest.approx(1_380_000_000)


def test_historicals_are_absolute_units(wb):
    ws = wb["02_Historical_Financials"]
    assert ws["B3"].value == 2022 and ws["E3"].value == 2025
    assert ws["B5"].value == pytest.approx(60_000_000 * 1000)  # 2022 revenue
    assert ws["E16"].value == pytest.approx(18_500_000 * 1000)  # 2025 cash
    # derived rows are formulas, not values
    assert ws["B7"].value == "=B5-B6"  # gross profit
    assert ws["E11"].value == "=E9-E10"  # EBIT = EBITDA - D&A


def test_dcf_is_fully_formula_driven(wb):
    ws = wb["04_DCF_Model"]
    assert ws["B5"].value == "='02_Historical_Financials'!$E$5"  # year 0 links to 02
    assert ws["C5"].value == "=B5*(1+Growth1)"
    assert ws["H5"].value == "=G5*(1+Growth2)"  # year 6 switches to stage 2
    assert ws["C6"].value == "=C5*EBITMargin"
    assert ws["C7"].value == "=C6*(1-TaxRate)"
    assert ws["C12"].value == '=IFERROR(1/(1+WACC)^C4,"")'
    tv = ws["B16"].value
    assert tv.startswith('=IF(OR(WACC="",WACC<=TerminalGrowth),""')
    assert "(1+TerminalGrowth)/(WACC-TerminalGrowth)" in tv
    # no hardcoded numbers anywhere in the forecast block (cols B..L, rows 5..13)
    for row in ws.iter_rows(min_row=5, max_row=13, min_col=2, max_col=12):
        for cell in row:
            if cell.value is not None:
                assert isinstance(cell.value, str) and cell.value.startswith("=")


def test_wacc_peer_clamp_and_blank_guards(wb):
    ws = wb["03_WACC_Calculation"]
    de = ws["K5"].value  # first peer D/E
    assert "MAX(0," in de and de.startswith("=IF(")
    unlev = ws["L5"].value
    assert "(1-TaxRate)" in unlev
    # median over the 4 peer rows
    assert ws["L9"].value == '=IFERROR(MEDIAN(L5:L8),"")'


def test_relative_valuation_references_sheet3(wb):
    ws = wb["05_Relative_Valuation"]
    assert "'03_WACC_Calculation'" in ws["B5"].value  # peer EV
    pe = ws["E5"].value
    assert "<=0" in pe  # negative-earnings P/E blanked


def test_summary_signal_and_weighting(wb, inputs):
    ws = wb["06_Valuation_Summary"]
    assert ws["B4"].value == "=DCFImpliedPrice"
    assert ws["C4"].value == "=WeightDCF"
    target = ws["B9"].value
    assert target.startswith("=IFERROR(SUMPRODUCT(")
    signal = ws["B12"].value
    assert '"AL"' in signal and '"SAT"' in signal and '"TUT"' in signal
    assert "SignalThreshold" in signal
    assert ws["B14"].value == '=IF(B9="","",IFERROR(B9/FXRate,""))'


def test_every_division_is_error_guarded(wb):
    """The spec's #DIV/0! safeguard: no bare division anywhere."""
    for name in SHEETS:
        for row in wb[name].iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("=") and "/" in v:
                    assert "IFERROR(" in v, f"{name}!{cell.coordinate}: unguarded division: {v}"


def test_english_locale_builds(inputs, tmp_path):
    from finauto.engine.builder import build_workbook

    en_inputs = inputs.model_copy(update={"locale": "en"})
    path = build_workbook(en_inputs, tmp_path / "en.xlsx")
    wb = openpyxl.load_workbook(path)
    signal = wb["06_Valuation_Summary"]["B12"].value
    assert '"BUY"' in signal and '"SELL"' in signal
